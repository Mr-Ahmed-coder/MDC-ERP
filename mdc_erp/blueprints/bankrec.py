"""Bank Reconciliation (Odoo Community 18-style)  (v8.0).

Import a bank statement (CSV or Excel), then reconcile each statement line
against the bank account's book journal lines:

  - automatic matching (exact amount + date window) and suggested matches
  - manual matching to a specific book entry
  - create a missing payment (posts a journal entry) for a line with no book
    counterpart, or a write-off (bank charges, small differences)
  - duplicate detection on import
  - reconciliation history, an approval workflow and full audit logging

Reconciled book lines are flagged via the existing JournalLine.cleared column,
so this ties into the ledger without changing it.
"""
import io
import csv
import datetime as dt
from flask import Blueprint, request, redirect, url_for, flash, abort
from markupsafe import escape as h
from ..extensions import db
from ..models import (BankAccount, BankStatement, BankStatementLine,
                      Account, JournalEntry, JournalLine)
from ..core.security import can, cur_user, log, login_required, csrf_token
from ..core.helpers import money, today
from ..core.ui import page
from ..core.posting import post_journal, acc

bp = Blueprint('bankrec', __name__)

MATCH_WINDOW_DAYS = 5


# --------------------------------------------------------------- helpers
def _approver():
    u = cur_user()
    return bool(u) and u.role in ('super_admin', 'accountant')


def _bank_jlines(bank, unmatched_only=True):
    """Book journal lines posted to this bank's COA account."""
    if not bank or not bank.account_id:
        return []
    q = (JournalLine.query.join(JournalEntry, JournalLine.entry_id == JournalEntry.id)
         .filter(JournalLine.account_id == bank.account_id))
    if unmatched_only:
        q = q.filter((JournalLine.cleared == False) | (JournalLine.cleared.is_(None)))  # noqa: E712
    return q.order_by(JournalEntry.date, JournalLine.id).all()


def _within(d1, d2, days=MATCH_WINDOW_DAYS):
    try:
        a = dt.date.fromisoformat((d1 or '')[:10])
        b = dt.date.fromisoformat((d2 or '')[:10])
        return abs((a - b).days) <= days
    except Exception:
        return True   # if dates unparseable, don't exclude on date


def _candidate(line, bank):
    """Best book journal line matching a statement line (amount + date)."""
    amt = line.amount or 0
    best = None
    for l in _bank_jlines(bank):
        booked = (l.debit or 0) - (l.credit or 0)      # +into bank / -out of bank
        if abs(booked - amt) < 0.005:
            if _within(line.date, l.entry.date):
                return l                                # strong match: amount + date
            best = best or l                            # amount matches, date off
    return best


# --------------------------------------------------------------- import parsing
def _num(s):
    s = str(s or '').replace(',', '').replace('$', '').strip()
    if not s:
        return 0.0
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()')
    try:
        v = float(s)
    except ValueError:
        return 0.0
    return -v if neg else v


def _pick(headers, *names):
    low = [h_.strip().lower() for h_ in headers]
    for n in names:
        if n in low:
            return low.index(n)
    for i, hh in enumerate(low):
        if any(n in hh for n in names):
            return i
    return None


def _parse_rows(rows):
    """rows = list of lists (first row headers). Returns list of dict lines."""
    if not rows:
        return []
    headers = [str(x) for x in rows[0]]
    di = _pick(headers, 'date')
    li = _pick(headers, 'label', 'description', 'details', 'narration', 'memo')
    ri = _pick(headers, 'ref', 'reference')
    ai = _pick(headers, 'amount', 'value')
    dbi = _pick(headers, 'debit', 'withdrawal', 'paid out')
    cri = _pick(headers, 'credit', 'deposit', 'paid in')
    out = []
    for r in rows[1:]:
        if not any(str(c).strip() for c in r):
            continue
        def cell(i):
            return r[i] if (i is not None and i < len(r)) else ''
        if ai is not None:
            amt = _num(cell(ai))
        else:
            amt = _num(cell(cri)) - _num(cell(dbi))   # credit positive, debit negative
        out.append(dict(date=str(cell(di))[:12], label=str(cell(li))[:200],
                        ref=str(cell(ri))[:80], amount=amt))
    return out


def _read_upload(file, pasted):
    name = (getattr(file, 'filename', '') or '').lower()
    if file and name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True)
        ws = wb.active
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        return _parse_rows(rows), name
    # CSV (upload or pasted text)
    text = ''
    if file and name.endswith('.csv'):
        text = file.read().decode('utf-8', 'ignore')
    elif pasted:
        text = pasted
    if not text:
        return [], name
    rows = list(csv.reader(io.StringIO(text)))
    return _parse_rows(rows), (name or 'pasted.csv')


# --------------------------------------------------------------- list / history
def bankrec_list():
    stmts = BankStatement.query.order_by(BankStatement.id.desc()).limit(100).all()
    rows = ''
    for s in stmts:
        n = len(s.lines)
        m = sum(1 for l in s.lines if l.matched)
        diff = (s.closing or 0) - (s.opening or 0) - sum(l.amount or 0 for l in s.lines)
        sc = {'Draft': 'amber', 'Reconciled': 'blue', 'Approved': 'green'}.get(s.status, 'grey')
        rows += (f"<tr><td><a href='{url_for('bankrec.reconcile', sid=s.id)}' style='color:var(--petrol);font-weight:600'>{h(s.name or ('Statement #%d' % s.id))}</a></td>"
                 f"<td>{h(s.bank_account.name if s.bank_account else '—')}</td><td>{h(s.date or '')}</td>"
                 f"<td class='num'>{m}/{n}</td>"
                 f"<td><div style='background:var(--canvas);border-radius:6px;height:8px;width:90px;overflow:hidden'>"
                 f"<div style='background:var(--green);height:8px;width:{int(100*m/n) if n else 0}%'></div></div></td>"
                 f"<td class='num'>{money(diff)}</td>"
                 f"<td><span class='pill {sc}'>{h(s.status)}</span></td>"
                 f"<td class='num'><a class='btn gh sm' href='{url_for('bankrec.reconcile', sid=s.id)}'>Open</a></td></tr>")
    if not rows:
        rows = "<tr><td colspan='8'><div class='empty'><b>No statements yet</b>Import a bank statement to begin.</div></td></tr>"
    body = (f"<div class='panel'><div class='ph'><h2>Bank Reconciliation</h2>"
            f"<span class='so'>Statements &amp; history</span><div class='sp'></div>"
            f"<a class='btn primary' href='{url_for('bankrec.import_stmt')}'>⭳ Import Statement</a></div>"
            f"<div class='tw'><table><thead><tr><th>Statement</th><th>Bank</th><th>Date</th>"
            f"<th class='num'>Matched</th><th>Progress</th><th class='num'>Difference</th><th>Status</th><th></th></tr></thead>"
            f"<tbody>{rows}</tbody></table></div></div>")
    return page('Bank Reconciliation', body, 'bankrec')


# --------------------------------------------------------------- import
@bp.route('/bankrec/import', methods=['GET', 'POST'])
@login_required
def import_stmt():
    if not can('bankrec'):
        abort(403)
    banks = BankAccount.query.filter_by(active=True).all()
    if request.method == 'POST':
        bank = BankAccount.query.get(int(request.form.get('bank', 0) or 0))
        if not bank:
            flash('Choose a bank account', 'err')
            return redirect(url_for('bankrec.import_stmt'))
        file = request.files.get('file')
        pasted = request.form.get('pasted', '')
        lines, fname = _read_upload(file, pasted)
        if not lines:
            flash('No transactions found — check the file/columns', 'err')
            return redirect(url_for('bankrec.import_stmt'))
        st = BankStatement(bank_account_id=bank.id,
                           name=request.form.get('name') or f"{bank.name} · {today()}",
                           date=request.form.get('date') or today(),
                           filename=fname,
                           opening=_num(request.form.get('opening')),
                           closing=_num(request.form.get('closing')),
                           status='Draft', created_by=cur_user().username)
        db.session.add(st)
        db.session.flush()
        # duplicate detection: same bank + date + amount + ref already imported
        existing = set()
        for l in (BankStatementLine.query.join(BankStatement)
                  .filter(BankStatement.bank_account_id == bank.id).all()):
            existing.add((l.date, round(l.amount or 0, 2), (l.ref or '').lower()))
        dups = 0
        for d in lines:
            key = (d['date'], round(d['amount'], 2), (d['ref'] or '').lower())
            is_dup = key in existing
            if is_dup:
                dups += 1
            existing.add(key)
            db.session.add(BankStatementLine(statement_id=st.id, date=d['date'], label=d['label'],
                                             ref=d['ref'], amount=d['amount'], is_duplicate=is_dup))
        db.session.commit()
        log(f'Imported bank statement {st.name} ({len(lines)} lines, {dups} duplicates)',
            action_type='import', entity=f'BankStatement#{st.id}')
        flash(f'Imported {len(lines)} lines' + (f' · {dups} possible duplicates flagged' if dups else ''), 'ok')
        return redirect(url_for('bankrec.reconcile', sid=st.id))

    bopts = ''.join(f"<option value='{b.id}'>{h(b.name)} ({h(b.bank_name or '')})</option>" for b in banks)
    if not banks:
        body = (f"<div class='panel'><div class='pad'><div class='empty'><b>No bank accounts</b>"
                f"Create one in <a href='{url_for('modules.module', mod='banks')}' style='color:var(--petrol)'>Bank Accounts</a> "
                f"and link it to a COA account (1102/1103).</div></div></div>")
        return page('Import Statement', body, 'bankrec')
    body = (f"<div class='panel' style='max-width:720px'><div class='ph'><h2>Import Bank Statement</h2>"
            f"<div class='sp'></div><a class='btn gh sm' href='{url_for('modules.module', mod='bankrec')}'>← Back</a></div>"
            f"<form method='post' enctype='multipart/form-data' class='pad' style='display:flex;flex-direction:column;gap:12px'>"
            f"<input type='hidden' name='_csrf' value='{h(csrf_token())}'>"
            f"<label>Bank account<br><select name='bank' class='lb-input' style='width:100%'>{bopts}</select></label>"
            f"<div style='display:flex;gap:10px;flex-wrap:wrap'>"
            f"<label>Statement name<br><input name='name' class='lb-input' placeholder='e.g. Premier Bank · August'></label>"
            f"<label>Date<br><input type='date' name='date' value='{today()}' class='lb-input'></label>"
            f"<label>Opening<br><input name='opening' class='lb-input' style='width:110px' placeholder='0'></label>"
            f"<label>Closing<br><input name='closing' class='lb-input' style='width:110px' placeholder='0'></label></div>"
            f"<label>Upload CSV or Excel<br><input type='file' name='file' accept='.csv,.xlsx,.xlsm' class='lb-input' style='width:100%'></label>"
            f"<div style='color:var(--muted);font-size:12px'>…or paste CSV below. Columns detected automatically: "
            f"<b>Date, Label/Description, Ref, Amount</b> (or separate <b>Debit/Credit</b>).</div>"
            f"<textarea name='pasted' rows='5' class='lb-input' placeholder='Date,Description,Ref,Amount&#10;2026-08-03,Deposit EVC,TX1001,290'></textarea>"
            f"<div><button class='btn primary'>Import</button></div></form></div>")
    return page('Import Statement', body, 'bankrec')


# --------------------------------------------------------------- reconcile workspace
@bp.route('/bankrec/<int:sid>')
@login_required
def reconcile(sid):
    if not can('bankrec'):
        abort(403)
    st = BankStatement.query.get_or_404(sid)
    bank = st.bank_account
    n = len(st.lines)
    m = sum(1 for l in st.lines if l.matched)
    locked = (st.status == 'Approved')

    rows = ''
    for l in st.lines:
        cand = None if l.matched else _candidate(l, bank)
        dup = " <span class='pill red' title='Same date/amount/ref already imported'>duplicate?</span>" if l.is_duplicate else ''
        amt_c = 'var(--green)' if (l.amount or 0) >= 0 else 'var(--red)'
        if l.matched:
            status = (f"<span class='pill green'>✓ {h(l.match_type or 'matched')}</span>"
                      + (f"<br><small style='color:var(--muted)'>{h(l.match_note or '')}</small>" if l.match_note else ''))
            action = '' if locked else (f"<a class='btn gh sm' href='{url_for('bankrec.line_unmatch', lid=l.id)}'>Unmatch</a>")
        else:
            sug = ''
            if cand:
                sug = (f"<div style='font-size:12px;color:var(--muted)'>Suggested: {h(cand.entry.ref or ('JV-%d' % cand.entry_id))} "
                       f"· {h(cand.entry.date)} · {money((cand.debit or 0) - (cand.credit or 0))}</div>")
            status = ("<span class='pill amber'>unmatched</span>" + sug)
            if locked:
                action = ''
            else:
                action = "<div style='display:flex;gap:4px;flex-wrap:wrap;justify-content:flex-end'>"
                if cand:
                    action += f"<a class='btn sm primary' href='{url_for('bankrec.line_match', lid=l.id, jl=cand.id)}'>Match</a>"
                action += (f"<a class='btn gh sm' href='{url_for('bankrec.line_match_pick', lid=l.id)}'>Manual</a>"
                           f"<a class='btn gh sm' href='{url_for('bankrec.line_payment', lid=l.id)}'>+ Payment</a>"
                           f"<a class='btn gh sm' href='{url_for('bankrec.line_writeoff', lid=l.id)}'>Write-off</a></div>")
        rows += (f"<tr><td>{h(l.date)}</td><td>{h(l.label or '')}{dup}</td><td>{h(l.ref or '')}</td>"
                 f"<td class='num' style='color:{amt_c};font-weight:600'>{money(l.amount)}</td>"
                 f"<td>{status}</td><td class='num'>{action}</td></tr>")
    if not rows:
        rows = "<tr><td colspan='6' style='color:var(--muted)'>No lines</td></tr>"

    diff = (st.closing or 0) - (st.opening or 0) - sum(l.amount or 0 for l in st.lines)
    pct = int(100 * m / n) if n else 0
    sc = {'Draft': 'amber', 'Reconciled': 'blue', 'Approved': 'green'}.get(st.status, 'grey')

    tools = ''
    if not locked:
        tools += f"<a class='btn' href='{url_for('bankrec.automatch', sid=sid)}'>⚡ Auto-match all</a> "
    if st.status == 'Draft' and m == n and n:
        tools += f"<a class='btn' href='{url_for('bankrec.mark_reconciled', sid=sid)}'>Mark reconciled</a> "
    if st.status == 'Reconciled' and _approver():
        tools += f"<a class='btn primary' href='{url_for('bankrec.approve', sid=sid)}'>✓ Approve</a> "
    if st.status == 'Approved':
        tools += (f"<span class='pill green'>Approved by {h(st.approved_by or '')} · {h(st.approved_at or '')}</span> ")

    summary = (f"<div class='pad' style='display:flex;gap:20px;flex-wrap:wrap;align-items:center;border-bottom:1px solid var(--line)'>"
               f"<div>Opening <b>{money(st.opening)}</b></div><div>Closing <b>{money(st.closing)}</b></div>"
               f"<div>Statement total <b>{money(sum(l.amount or 0 for l in st.lines))}</b></div>"
               f"<div>Difference <b style='color:{'var(--green)' if abs(diff)<0.5 else 'var(--red)'}'>{money(diff)}</b></div>"
               f"<div style='min-width:160px'>Matched {m}/{n}<div style='background:var(--canvas);border-radius:6px;height:8px;margin-top:3px'>"
               f"<div style='background:var(--green);height:8px;border-radius:6px;width:{pct}%'></div></div></div>"
               f"<div style='margin-left:auto'>{tools}</div></div>")

    body = (f"<div class='panel'><div class='ph'><h2>{h(st.name or ('Statement #%d' % st.id))}</h2>"
            f"<span class='so'>{h(bank.name if bank else '')} · {h(st.date or '')}</span>"
            f"<span class='pill {sc}' style='margin-left:8px'>{h(st.status)}</span><div class='sp'></div>"
            f"<a class='btn gh sm' href='{url_for('modules.module', mod='bankrec')}'>← All statements</a></div>"
            + summary
            + f"<div class='tw'><table><thead><tr><th>Date</th><th>Label</th><th>Ref</th>"
            f"<th class='num'>Amount</th><th>Status</th><th></th></tr></thead><tbody>{rows}</tbody></table></div></div>")
    return page('Reconcile', body, 'bankrec')


# --------------------------------------------------------------- line actions
def _guard_line(lid):
    if not can('bankrec'):
        abort(403)
    l = BankStatementLine.query.get_or_404(lid)
    if l.statement.status == 'Approved':
        flash('Statement is approved and locked', 'err')
        return None, l
    return l, l


def _mark_cleared(jl, cleared=True):
    """Mark all lines of the journal entry cleared (whole move reconciled)."""
    if jl:
        for sib in JournalLine.query.filter_by(entry_id=jl.entry_id).all():
            sib.cleared = cleared


@bp.route('/bankrec/line/<int:lid>/match/<int:jl>')
@login_required
def line_match(lid, jl):
    l, _ = _guard_line(lid)
    if not l:
        return redirect(request.referrer or url_for('modules.module', mod='bankrec'))
    j = JournalLine.query.get_or_404(jl)
    l.matched = True
    l.match_type = 'manual' if request.args.get('m') else 'auto'
    l.match_note = f"{j.entry.ref or ('JV-%d' % j.entry_id)} · {j.entry.date}"
    l.journal_entry_id = j.entry_id
    _mark_cleared(j, True)
    db.session.commit()
    log(f'Matched bank line #{l.id} → {l.match_note}', action_type='reconcile', entity=f'BankStatementLine#{l.id}')
    return redirect(url_for('bankrec.reconcile', sid=l.statement_id))


@bp.route('/bankrec/line/<int:lid>/pick')
@login_required
def line_match_pick(lid):
    l, _ = _guard_line(lid)
    if not l:
        return redirect(url_for('modules.module', mod='bankrec'))
    bank = l.statement.bank_account
    opts = ''
    for j in _bank_jlines(bank):
        booked = (j.debit or 0) - (j.credit or 0)
        opts += (f"<tr><td>{h(j.entry.date)}</td><td>{h(j.entry.ref or ('JV-%d' % j.entry_id))}</td>"
                 f"<td>{h((j.entry.memo or '')[:50])}</td><td class='num'>{money(booked)}</td>"
                 f"<td class='num'><a class='btn sm primary' href='{url_for('bankrec.line_match', lid=l.id, jl=j.id, m=1)}'>Match</a></td></tr>")
    if not opts:
        opts = "<tr><td colspan='5' style='color:var(--muted)'>No unreconciled book entries on this bank account.</td></tr>"
    body = (f"<div class='panel'><div class='ph'><h2>Manual match</h2>"
            f"<span class='so'>{h(l.date)} · {h(l.label or '')} · {money(l.amount)}</span><div class='sp'></div>"
            f"<a class='btn gh sm' href='{url_for('bankrec.reconcile', sid=l.statement_id)}'>← Cancel</a></div>"
            f"<div class='tw'><table><thead><tr><th>Date</th><th>Move</th><th>Memo</th><th class='num'>Amount</th><th></th></tr></thead>"
            f"<tbody>{opts}</tbody></table></div></div>")
    return page('Manual match', body, 'bankrec')


def _counter_accounts():
    return Account.query.filter(Account.is_group == False, Account.active == True).order_by(Account.code).all()  # noqa: E712


@bp.route('/bankrec/line/<int:lid>/payment', methods=['GET', 'POST'])
@login_required
def line_payment(lid):
    l, _ = _guard_line(lid)
    if not l:
        return redirect(url_for('modules.module', mod='bankrec'))
    bank = l.statement.bank_account
    bank_acc = Account.query.get(bank.account_id) if bank and bank.account_id else None
    if request.method == 'POST':
        counter = acc(request.form.get('account', ''))
        if not (bank_acc and counter):
            flash('Pick a counterpart account', 'err')
            return redirect(url_for('bankrec.line_payment', lid=lid))
        amt = l.amount or 0
        ref = f"BANK-{l.statement_id}-{l.id}"
        if amt >= 0:      # money into bank
            lines = [(bank_acc.code, amt, 0), (counter.code, 0, amt)]
        else:             # money out of bank
            lines = [(counter.code, -amt, 0), (bank_acc.code, 0, -amt)]
        je = post_journal(l.date or today(), ref, f"Bank payment: {l.label or ''}"[:200], lines)
        if je:
            for sib in je.lines:
                sib.cleared = True
            l.matched = True
            l.match_type = 'payment'
            l.match_note = f"Created {ref} → {counter.code} {counter.name}"
            l.journal_entry_id = je.id
            db.session.commit()
            log(f'Created payment for bank line #{l.id} ({ref})', action_type='reconcile', entity=f'BankStatementLine#{l.id}')
            flash('Payment created and matched', 'ok')
        return redirect(url_for('bankrec.reconcile', sid=l.statement_id))
    aopts = ''.join(f"<option value='{a.code}'>{h(a.code or '')} · {h(a.name)}</option>" for a in _counter_accounts())
    body = (f"<div class='panel' style='max-width:560px'><div class='ph'><h2>Create missing payment</h2>"
            f"<div class='sp'></div><a class='btn gh sm' href='{url_for('bankrec.reconcile', sid=l.statement_id)}'>← Cancel</a></div>"
            f"<form method='post' class='pad' style='display:flex;flex-direction:column;gap:12px'>"
            f"<input type='hidden' name='_csrf' value='{h(csrf_token())}'>"
            f"<div>Line: <b>{h(l.date)} · {h(l.label or '')}</b> · <b style='color:{'var(--green)' if (l.amount or 0)>=0 else 'var(--red)'}'>{money(l.amount)}</b></div>"
            f"<div style='color:var(--muted);font-size:13px'>Posts a balanced entry to <b>{h(bank_acc.code if bank_acc else '—')} {h(bank_acc.name if bank_acc else '')}</b> and the counterpart account you choose, then marks the line reconciled.</div>"
            f"<label>Counterpart account<br><select name='account' class='lb-input' style='width:100%'>{aopts}</select></label>"
            f"<div><button class='btn primary'>Create &amp; match</button></div></form></div>")
    return page('Create payment', body, 'bankrec')


@bp.route('/bankrec/line/<int:lid>/writeoff', methods=['GET', 'POST'])
@login_required
def line_writeoff(lid):
    l, _ = _guard_line(lid)
    if not l:
        return redirect(url_for('modules.module', mod='bankrec'))
    bank = l.statement.bank_account
    bank_acc = Account.query.get(bank.account_id) if bank and bank.account_id else None
    if request.method == 'POST':
        counter = acc(request.form.get('account', ''))
        amt = _num(request.form.get('amount')) or (l.amount or 0)
        if not (bank_acc and counter):
            flash('Pick a write-off account', 'err')
            return redirect(url_for('bankrec.line_writeoff', lid=lid))
        ref = f"WOFF-{l.statement_id}-{l.id}"
        if amt >= 0:
            lines = [(bank_acc.code, amt, 0), (counter.code, 0, amt)]
        else:
            lines = [(counter.code, -amt, 0), (bank_acc.code, 0, -amt)]
        je = post_journal(l.date or today(), ref, f"Write-off: {l.label or ''}"[:200], lines)
        if je:
            for sib in je.lines:
                sib.cleared = True
            l.matched = True
            l.match_type = 'writeoff'
            l.match_note = f"Write-off {money(amt)} → {counter.code}"
            l.journal_entry_id = je.id
            db.session.commit()
            log(f'Write-off for bank line #{l.id} ({ref})', action_type='reconcile', entity=f'BankStatementLine#{l.id}')
            flash('Write-off posted and matched', 'ok')
        return redirect(url_for('bankrec.reconcile', sid=l.statement_id))
    aopts = ''.join(f"<option value='{a.code}'>{h(a.code or '')} · {h(a.name)}</option>"
                    for a in _counter_accounts() if a.type in ('Expense', 'Income'))
    body = (f"<div class='panel' style='max-width:560px'><div class='ph'><h2>Write-off</h2>"
            f"<div class='sp'></div><a class='btn gh sm' href='{url_for('bankrec.reconcile', sid=l.statement_id)}'>← Cancel</a></div>"
            f"<form method='post' class='pad' style='display:flex;flex-direction:column;gap:12px'>"
            f"<input type='hidden' name='_csrf' value='{h(csrf_token())}'>"
            f"<div>Line: <b>{h(l.label or '')}</b> · {money(l.amount)}</div>"
            f"<div style='color:var(--muted);font-size:13px'>For bank charges, fees or small differences. Posts to the bank account and an expense/income account.</div>"
            f"<label>Amount<br><input name='amount' class='lb-input' value='{l.amount or 0}'></label>"
            f"<label>Write-off account<br><select name='account' class='lb-input' style='width:100%'>{aopts}</select></label>"
            f"<div><button class='btn primary'>Post write-off</button></div></form></div>")
    return page('Write-off', body, 'bankrec')


@bp.route('/bankrec/line/<int:lid>/unmatch')
@login_required
def line_unmatch(lid):
    l, _ = _guard_line(lid)
    if not l:
        return redirect(url_for('modules.module', mod='bankrec'))
    if l.journal_entry_id:
        for sib in JournalLine.query.filter_by(entry_id=l.journal_entry_id).all():
            sib.cleared = False
    l.matched = False
    l.match_type = None
    l.match_note = None
    l.journal_entry_id = None
    db.session.commit()
    log(f'Unmatched bank line #{l.id}', action_type='reconcile', entity=f'BankStatementLine#{l.id}')
    return redirect(url_for('bankrec.reconcile', sid=l.statement_id))


# --------------------------------------------------------------- statement actions
@bp.route('/bankrec/<int:sid>/automatch')
@login_required
def automatch(sid):
    if not can('bankrec'):
        abort(403)
    st = BankStatement.query.get_or_404(sid)
    if st.status == 'Approved':
        flash('Statement is approved and locked', 'err')
        return redirect(url_for('bankrec.reconcile', sid=sid))
    bank = st.bank_account
    n = 0
    for l in st.lines:
        if l.matched:
            continue
        cand = _candidate(l, bank)
        if cand and _within(l.date, cand.entry.date):
            l.matched = True
            l.match_type = 'auto'
            l.match_note = f"{cand.entry.ref or ('JV-%d' % cand.entry_id)} · {cand.entry.date}"
            l.journal_entry_id = cand.entry_id
            _mark_cleared(cand, True)
            n += 1
    db.session.commit()
    log(f'Auto-matched {n} lines on statement #{sid}', action_type='reconcile', entity=f'BankStatement#{sid}')
    flash(f'Auto-matched {n} line(s)', 'ok')
    return redirect(url_for('bankrec.reconcile', sid=sid))


@bp.route('/bankrec/<int:sid>/reconciled')
@login_required
def mark_reconciled(sid):
    if not can('bankrec'):
        abort(403)
    st = BankStatement.query.get_or_404(sid)
    if any(not l.matched for l in st.lines):
        flash('All lines must be matched first', 'err')
        return redirect(url_for('bankrec.reconcile', sid=sid))
    st.status = 'Reconciled'
    db.session.commit()
    log(f'Statement #{sid} marked reconciled', action_type='reconcile', entity=f'BankStatement#{sid}')
    flash('Marked reconciled — awaiting approval', 'ok')
    return redirect(url_for('bankrec.reconcile', sid=sid))


@bp.route('/bankrec/<int:sid>/approve')
@login_required
def approve(sid):
    if not _approver():
        abort(403)
    st = BankStatement.query.get_or_404(sid)
    if st.status != 'Reconciled':
        flash('Only a reconciled statement can be approved', 'err')
        return redirect(url_for('bankrec.reconcile', sid=sid))
    st.status = 'Approved'
    st.approved_by = cur_user().username
    st.approved_at = today()
    db.session.commit()
    log(f'Statement #{sid} APPROVED', action_type='approve', entity=f'BankStatement#{sid}')
    flash('Statement approved and locked', 'ok')
    return redirect(url_for('bankrec.reconcile', sid=sid))
