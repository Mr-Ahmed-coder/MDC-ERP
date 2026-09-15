"""Financial exports: Excel workbooks (openpyxl) and PDF statements pack (reportlab)."""
import datetime as dt
import io

from flask import Blueprint, request, send_file, abort
from ..models import (Account, JournalEntry, JournalLine, Invoice, Purchase,
                      Expense, Budget)
from ..core.security import cur_user, login_required, can

bp = Blueprint('finexport', __name__)

BRAND_BLUE = '1A3E8F'
BRAND_ORANGE = 'F57C00'


def _wb():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill('solid', fgColor=BRAND_BLUE)
    return wb, hdr_font, hdr_fill, Alignment


def _sheet(ws, title, headers, rows, hdr_font, hdr_fill, money_cols=()):
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = hdr_font
        c.fill = hdr_fill
    for r in rows:
        ws.append(r)
    for idx in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for c in row:
                c.number_format = '#,##0.00'
    for col in ws.columns:
        width = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(width + 3, 46)


def _send_wb(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name=name, as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _account_balances(year=None):
    """code -> (account, debit_total, credit_total) optionally limited to a year.

    Opening balances from the Chart of Accounts are folded in (debit-positive
    by account type) so opening-only accounts appear and TB/BS tie to the COA.
    """
    from ..core.posting import opening_dr
    out = {}
    for a in Account.query.all():
        od = opening_dr(a)
        if abs(od) > 0.005 and a.code:
            out[a.code] = (a, od if od > 0 else 0.0, -od if od < 0 else 0.0)
    q = JournalLine.query.join(JournalEntry)
    if year:
        q = q.filter(JournalEntry.date.like(f'{year}%'))
    for l in q.all():
        a = l.account
        if not a:
            continue
        d, c = out.get(a.code, (a, 0.0, 0.0))[1:] if a.code in out else (0.0, 0.0)
        out[a.code] = (a, d + (l.debit or 0), c + (l.credit or 0))
    return out


@bp.route('/export/xlsx/<what>')
@login_required
def export_xlsx(what):
    y = request.args.get('year', type=int) or dt.date.today().year
    wb, hf, hfill, _ = _wb()
    ws = wb.active
    if what == 'trialbalance':
        if not can('acct'): abort(403)
        rows = []
        for code, (a, d, c) in sorted(_account_balances(y).items()):
            net = d - c
            rows.append([code, a.name, a.type, round(d, 2), round(c, 2),
                         round(net, 2) if net > 0 else 0, round(-net, 2) if net < 0 else 0])
        _sheet(ws, f'Trial Balance {y}',
               ['Code', 'Account', 'Type', 'Total Debit', 'Total Credit', 'Dr Balance', 'Cr Balance'],
               rows, hf, hfill, money_cols=(4, 5, 6, 7))
        return _send_wb(wb, f'trial-balance-{y}.xlsx')
    if what == 'ledger':
        if not can('ledger'): abort(403)
        rows = []
        for l in (JournalLine.query.join(JournalEntry)
                  .filter(JournalEntry.date.like(f'{y}%'))
                  .order_by(JournalEntry.date, JournalEntry.id).all()):
            e = l.entry
            a = l.account
            rows.append([e.date if e else '', e.ref if e else '',
                         f'{a.code} · {a.name}' if a else '', (e.memo or '') if e else '',
                         round(l.debit or 0, 2), round(l.credit or 0, 2),
                         'Yes' if l.cleared else ''])
        _sheet(ws, f'General Ledger {y}',
               ['Date', 'Ref', 'Account', 'Memo', 'Debit', 'Credit', 'Cleared'],
               rows, hf, hfill, money_cols=(5, 6))
        return _send_wb(wb, f'general-ledger-{y}.xlsx')
    if what == 'araging':
        if not can('araging'): abort(403)
        rows = [[(i.patient.name if i.patient else 'Walk-in'), i.date, f'INV-{i.id:04d}',
                 (dt.date.today() - dt.date.fromisoformat(i.date)).days if i.date else 0,
                 round(i.balance, 2)]
                for i in Invoice.query.all()
                if i.status != 'Cancelled' and i.balance > 0.005]
        _sheet(ws, 'AR Aging', ['Patient', 'Date', 'Invoice', 'Age (days)', 'Balance'],
               rows, hf, hfill, money_cols=(5,))
        return _send_wb(wb, 'ar-aging.xlsx')
    if what == 'apaging':
        if not can('apaging'): abort(403)
        rows = []
        for p in Purchase.query.all():
            bal = (p.total or 0) - (p.paid or 0)
            if bal > 0.005:
                rows.append([(p.supplier.name if p.supplier else (p.item or '—')),
                             p.date, f'PUR-{p.id:04d}', round(bal, 2)])
        for e in Expense.query.all():
            bal = (e.amount or 0) - (e.paid or 0)
            if bal > 0.005:
                rows.append([f'{e.category or "Expense"}', e.date, f'EXP-{e.id:04d}', round(bal, 2)])
        _sheet(ws, 'AP Aging', ['Supplier / Item', 'Date', 'Ref', 'Balance'],
               rows, hf, hfill, money_cols=(4,))
        return _send_wb(wb, 'ap-aging.xlsx')
    if what == 'budget':
        if not can('budgetreport'): abort(403)
        bal = _account_balances(y)
        rows = []
        for b in Budget.query.filter_by(year=y).all():
            a = b.account
            act = 0.0
            if a and a.code in bal:
                _, d, c = bal[a.code]
                act = (c - d) if a.type == 'Income' else (d - c)
            rows.append([f'{a.code} · {a.name}' if a else '—', a.type if a else '',
                         round(b.amount or 0, 2), round(act, 2),
                         round((b.amount or 0) - act, 2)])
        _sheet(ws, f'Budget {y}', ['Account', 'Type', 'Budget', 'Actual', 'Remaining'],
               rows, hf, hfill, money_cols=(3, 4, 5))
        return _send_wb(wb, f'budget-{y}.xlsx')
    abort(404)


@bp.route('/finance/pdf')
@login_required
def finance_pdf():
    """Financial statements pack: Income Statement + Balance Sheet + Trial Balance."""
    if not can('acct'):
        abort(403)
    y = request.args.get('year', type=int) or dt.date.today().year
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, PageBreak)
    from ..core.security import setting

    blue = colors.HexColor('#' + BRAND_BLUE)
    orange = colors.HexColor('#' + BRAND_ORANGE)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle('H1', parent=styles['Title'], textColor=blue, fontSize=17, spaceAfter=2)
    h2 = ParagraphStyle('H2', parent=styles['Heading2'], textColor=blue, spaceBefore=10)
    small = ParagraphStyle('S', parent=styles['Normal'], fontSize=8.5, textColor=colors.grey)

    def tbl(data, money_from=1):
        t = Table(data, hAlign='LEFT', colWidths=None)
        style = [('BACKGROUND', (0, 0), (-1, 0), blue),
                 ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                 ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                 ('FONTSIZE', (0, 0), (-1, -1), 8.5),
                 ('ALIGN', (money_from, 1), (-1, -1), 'RIGHT'),
                 ('LINEBELOW', (0, 0), (-1, -2), 0.25, colors.HexColor('#DDDDDD')),
                 ('LINEABOVE', (0, -1), (-1, -1), 1, orange),
                 ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                 ('TOPPADDING', (0, 0), (-1, -1), 4),
                 ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]
        t.setStyle(TableStyle(style))
        return t

    m = lambda v: f'{v:,.2f}'
    bal = _account_balances(y)
    inc = [(a, c - d) for a, d, c in bal.values() if a.type == 'Income' and abs(c - d) > 0.005]
    exp = [(a, d - c) for a, d, c in bal.values() if a.type == 'Expense' and abs(d - c) > 0.005]
    rev_t = sum(v for _, v in inc)
    exp_t = sum(v for _, v in exp)
    net = rev_t - exp_t

    # balance sheet (all-time balances)
    bal_all = _account_balances(None)
    def side(t):
        sign = 1 if t == 'Asset' else -1
        return [(a, sign * (d - c)) for a, d, c in bal_all.values()
                if a.type == t and abs(d - c) > 0.005]
    assets = side('Asset')          # openings already folded into _account_balances
    liabs = side('Liability')
    equity = side('Equity')
    from ..core.posting import opening_totals
    _, _, _, _obe = opening_totals()
    if abs(_obe) > 0.005:           # standard offset so openings self-balance
        equity = equity + [(type('OBE', (), {'code': '3199', 'name': 'Opening Balance Equity (auto)'})(), _obe)]
    a_t = sum(v for _, v in assets)
    l_t = sum(v for _, v in liabs)
    e_t = sum(v for _, v in equity) + net

    company = setting('company', 'Modern Diagnostic Center')
    _addr = setting('company_address', '') or ''
    _phone = setting('company_phone', '') or ''
    _email = setting('company_email', '') or ''
    _contact = ' · '.join(x for x in [_addr, _phone, _email] if x)

    def _decorate(canv, docu):
        """Branded header band + footer drawn on every page."""
        canv.saveState()
        W, Hh = A4
        # header band
        canv.setFillColor(blue)
        canv.rect(0, Hh - 15 * mm, W, 15 * mm, fill=1, stroke=0)
        canv.setStrokeColor(orange); canv.setLineWidth(1.4)
        canv.line(0, Hh - 15 * mm, W, Hh - 15 * mm)
        canv.setFillColor(colors.white)
        canv.setFont('Helvetica-Bold', 12.5)
        canv.drawString(15 * mm, Hh - 10 * mm, company)
        if _contact:
            canv.setFont('Helvetica', 7.5)
            canv.drawRightString(W - 15 * mm, Hh - 9.5 * mm, _contact[:96])
        # footer
        canv.setStrokeColor(colors.HexColor('#CCCCCC')); canv.setLineWidth(0.5)
        canv.line(15 * mm, 13 * mm, W - 15 * mm, 13 * mm)
        canv.setFillColor(colors.grey); canv.setFont('Helvetica', 7.5)
        canv.drawString(15 * mm, 8.5 * mm, (company + ('  ·  ' + _contact if _contact else ''))[:118])
        canv.drawRightString(W - 15 * mm, 8.5 * mm, f'Page {docu.page}')
        canv.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=22 * mm, bottomMargin=18 * mm)
    story = [Paragraph(f'Financial Statements &mdash; FY {y}', h1),
             Paragraph(f'Fiscal Year {y} · '
                       f'generated {dt.date.today().isoformat()} by '
                       f'{cur_user().name or cur_user().username}', small),
             Spacer(1, 6 * mm),
             Paragraph('Income Statement', h2),
             tbl([['Account', f'FY {y}']]
                 + [[f'{a.code} · {a.name}', m(v)] for a, v in sorted(inc, key=lambda x: x[0].code)]
                 + [['Total Revenue', m(rev_t)]]),
             Spacer(1, 3 * mm),
             tbl([['Expenses', '']]
                 + [[f'{a.code} · {a.name}', m(v)] for a, v in sorted(exp, key=lambda x: x[0].code)]
                 + [['Total Expenses', m(exp_t)]]),
             Spacer(1, 3 * mm),
             tbl([['', ''], ['NET INCOME', m(net)]]),
             PageBreak(),
             Paragraph('Balance Sheet', h2),
             tbl([['Assets', '']]
                 + [[f'{a.code} · {a.name}', m(v)] for a, v in sorted(assets, key=lambda x: x[0].code)]
                 + [['Total Assets', m(a_t)]]),
             Spacer(1, 3 * mm),
             tbl([['Liabilities & Equity', '']]
                 + [[f'{a.code} · {a.name}', m(v)] for a, v in sorted(liabs + equity, key=lambda x: x[0].code)]
                 + [['Net Income (current)', m(net)],
                    ['Total Liabilities & Equity', m(l_t + e_t)]]),
             PageBreak(),
             Paragraph(f'Trial Balance · FY {y}', h2),
             tbl([['Code', 'Account', 'Debit', 'Credit']]
                 + [[a.code, a.name, m(d), m(c)]
                    for _, (a, d, c) in sorted(bal.items())]
                 + [['', 'TOTALS',
                     m(sum(d for _, d, _c in bal.values())),
                     m(sum(c for _, _d, c in bal.values()))]], money_from=2)]
    doc.build(story, onFirstPage=_decorate, onLaterPages=_decorate)
    buf.seek(0)
    return send_file(buf, download_name=f'financial-statements-{y}.pdf',
                     as_attachment=True, mimetype='application/pdf')


# ---- CSV exports (mirror of the xlsx reports; accountants often prefer CSV) ----
@bp.route('/export/csv/<what>')
@login_required
def export_csv(what):
    import csv
    from flask import Response
    y = request.args.get('year', type=int) or dt.date.today().year
    out = io.StringIO()
    w = csv.writer(out)

    if what == 'trialbalance':
        if not can('acct'): abort(403)
        w.writerow(['Code', 'Account', 'Type', 'Total Debit', 'Total Credit', 'Dr Balance', 'Cr Balance'])
        for code, (a, d, c) in sorted(_account_balances(y).items()):
            net = d - c
            w.writerow([code, a.name, a.type, round(d, 2), round(c, 2),
                        round(net, 2) if net > 0 else 0, round(-net, 2) if net < 0 else 0])
        fn = f'trial-balance-{y}.csv'
    elif what == 'ledger':
        if not can('ledger'): abort(403)
        w.writerow(['Date', 'Ref', 'Account', 'Memo', 'Debit', 'Credit', 'Cleared'])
        for l in (JournalLine.query.join(JournalEntry)
                  .filter(JournalEntry.date.like(f'{y}%'))
                  .order_by(JournalEntry.date, JournalEntry.id).all()):
            e = l.entry; a = l.account
            w.writerow([e.date if e else '', e.ref if e else '',
                        f'{a.code} · {a.name}' if a else '', (e.memo or '') if e else '',
                        round(l.debit or 0, 2), round(l.credit or 0, 2), 'Yes' if l.cleared else ''])
        fn = f'general-ledger-{y}.csv'
    elif what == 'araging':
        if not can('araging'): abort(403)
        w.writerow(['Patient', 'Date', 'Invoice', 'Age (days)', 'Balance'])
        for i in Invoice.query.all():
            if i.status != 'Cancelled' and i.balance > 0.005:
                age = (dt.date.today() - dt.date.fromisoformat(i.date)).days if i.date else 0
                w.writerow([(i.patient.name if i.patient else 'Walk-in'), i.date,
                            f'INV-{i.id:04d}', age, round(i.balance, 2)])
        fn = 'ar-aging.csv'
    elif what == 'invoices':
        if not can('invoices'): abort(403)
        w.writerow(['Invoice', 'Date', 'Patient', 'Subtotal', 'Discount', 'VAT', 'Total', 'Paid', 'Balance', 'Status'])
        for i in Invoice.query.order_by(Invoice.id).all():
            w.writerow([f'INV-{i.id:04d}', i.date, (i.patient.name if i.patient else 'Walk-in'),
                        round(i.subtotal, 2), round(i.discount or 0, 2), round(i.vat or 0, 2),
                        round(i.total, 2), round(i.paid or 0, 2), round(i.balance, 2), i.status or ''])
        fn = 'invoices.csv'
    else:
        abort(404)

    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': f'attachment;filename={fn}'})
